from openpyxl import load_workbook
from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, time, timedelta
import json, hashlib, re, sys

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'source' / 'PrepLadder_Subject_Checklist.xlsx'
OUT = ROOT / 'src' / 'seed.json'
OVERRIDES = {
    ('Forensic Medicine', 7): 39*60,
    ('Forensic Medicine', 15): 29*60,
    ('Forensic Medicine', 34): 71*60,
    ('Internal Medicine', 77): 9*60,
    ('Internal Medicine', 94): 36*60,
    ('Surgery', 3): 13*60,
    ('Surgery', 43): 71*60,
    ('PSM', 1): 43*60,
    ('Dermatology', 3): 60*60,
}
wb = load_workbook(SOURCE, read_only=True, data_only=True)
subjects, lectures, issues, seen = [], [], [], set()

def seconds(value):
    if value is None or str(value).strip() == '': return None
    if isinstance(value, timedelta): return int(value.total_seconds())
    if isinstance(value, datetime): return value.hour*3600+value.minute*60+value.second
    if isinstance(value, time): return value.hour*3600+value.minute*60+value.second
    if isinstance(value, (int,float)) and not isinstance(value,bool):
        return int(round(value*86400)) if 0 <= value < 1 else None
    s=str(value).strip().lower()
    m=re.fullmatch(r'(\d{1,3}):(\d{2})(?::(\d{2}))?',s)
    if m:
        a,b,c=[int(x or 0) for x in m.groups()]
        return a*3600+b*60+c if b<60 and c<60 else None
    m=re.fullmatch(r'(?:(\d+(?:\.\d+)?)\s*(?:h|hr|hrs|hour|hours)\s*)?(?:(\d+(?:\.\d+)?)\s*(?:m|min|mins|minute|minutes))?',s)
    if m and (m.group(1) or m.group(2)):
        return round(float(m.group(1) or 0)*3600+float(m.group(2) or 0)*60)
    return None

for wi, ws in enumerate(wb.worksheets):
    rows=list(ws.iter_rows(values_only=True))
    if tuple(str(x or '').strip().lower() for x in rows[0][:4]) != ('no.','lecture title','duration (h:mm)','completed'):
        raise ValueError(f'Unexpected header in {ws.title}: {rows[0][:4]}')
    subject_id=hashlib.sha256(ws.title.encode('utf-8')).hexdigest()[:16]
    subjects.append({'id':subject_id,'name':ws.title,'order':wi})
    numbers=[]
    for ri,row in enumerate(rows[1:],start=2):
        if not any(v is not None for v in row): continue
        number,title,raw_duration,status=(list(row)+[None]*4)[:4]
        if not isinstance(title,str) or not title.strip():
            issues.append({'sheet':ws.title,'row':ri,'issue':'missing title','raw':str(row[:4])});continue
        title=title.strip()
        duration=seconds(raw_duration)
        supplied_duration=False
        if duration is None and str(number).strip().isdigit() and (ws.title,int(number)) in OVERRIDES:
            duration=OVERRIDES[(ws.title,int(number))]
            supplied_duration=True
        if duration is None:
            issues.append({'sheet':ws.title,'row':ri,'lecture':title,'issue':'missing or unparseable duration','raw':str(raw_duration)})
        elif duration<=0:
            issues.append({'sheet':ws.title,'row':ri,'lecture':title,'issue':'non-positive duration','raw':str(raw_duration)})
            duration=None
        if status not in ('☑','☐',None,''):
            issues.append({'sheet':ws.title,'row':ri,'lecture':title,'issue':'unknown completion status','raw':str(status)})
        number=str(number).strip() if number is not None else ''
        if number.isdigit(): numbers.append(int(number))
        key=f'{ws.title}\u001f{ri}\u001f{title}'
        lid=hashlib.sha256(key.encode('utf-8')).hexdigest()[:24]
        if lid in seen: raise ValueError(f'Duplicate id: {key}')
        seen.add(lid)
        lectures.append({'id':lid,'subjectId':subject_id,'number':number,'title':title,'durationSeconds':duration,'chapter':None,'order':len(lectures),'sourceRow':ri,'sourceSheet':ws.title,'durationSource':'user-supplied' if supplied_duration else 'workbook','initialCompleted':status=='☑'})
    if len(numbers)!=len(set(numbers)):
        issues.append({'sheet':ws.title,'issue':'duplicate lecture number','numbers':[n for n,c in Counter(numbers).items() if c>1]})

course={'schemaVersion':1,'courseId':'prepladder-subject-checklist-2026-10-17','sourceFile':SOURCE.name,'subjects':subjects,'lectures':lectures}
OUT.parent.mkdir(exist_ok=True)
OUT.write_text(json.dumps(course,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
summary=[]
for s in subjects:
    items=[l for l in lectures if l['subjectId']==s['id']]
    summary.append({'subject':s['name'],'lectures':len(items),'knownHours':round(sum(l['durationSeconds'] or 0 for l in items)/3600,2),'unknown':sum(l['durationSeconds'] is None for l in items),'initialCompleted':sum(l['initialCompleted'] for l in items)})
report={'source':SOURCE.name,'subjects':summary,'totalLectures':len(lectures),'knownSeconds':sum(l['durationSeconds'] or 0 for l in lectures),'unknownDurations':sum(l['durationSeconds'] is None for l in lectures),'initialCompleted':sum(l['initialCompleted'] for l in lectures),'issues':issues}
(ROOT / 'source' / 'audit-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
